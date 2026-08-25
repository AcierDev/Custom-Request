from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document


LAST_UPDATED_OLD = "Last updated: August 3, 2026"
LAST_UPDATED_NEW = "Last updated: August 4, 2026"

PARAGRAPH_REPLACEMENTS = {
    LAST_UPDATED_OLD: LAST_UPDATED_NEW,
    "The account labeled US Bank in Zoho Books is the company's Amazon credit card.": (
        "The U.S. Bank account referenced in prior notes is the company's Amazon credit card. "
        "In Zoho Books it is named Amazon Business CC and classified as a Credit Card, not a bank account."
    ),
    (
        "The currently selected presentation is to show statement credits from older rewards points as Other Income, "
        "not sales revenue. This increases book profit in the period of the journal entry. The rewards were earned in "
        "previous years, so their tax treatment and period attribution should be confirmed with the tax preparer."
    ): (
        "Statement credits from older rewards points are presented in the Credit Card Rewards account, which is "
        "classified as Other Income. The March and April journal entries reclassified amounts from General Income to "
        "Credit Card Rewards and did not change total profit. The rewards were earned in previous years, so their tax "
        "treatment and period attribution should be confirmed with the tax preparer."
    ),
    "Shipping-label reimbursements": "Shipping-related corrections",
    (
        "The $28.13 and $700 Zelle payments involving Benjamin relate to Pirate Ship labels. Use this decision rule:"
    ): (
        "The $28.13 transaction dated July 1, 2026 is a FedEx refund credited to Postage. It is unrelated to "
        "Benjamin Steele, Yaakov Weil, or Pirate Ship."
    ),
    "If the Pirate Ship charges are not already recorded in Zoho Books, classify the reimbursement as Postage.": (
        "The $700.78 Zelle payment dated July 14, 2026 reimbursed Yaakov Weil for Pirate Ship labels accidentally "
        "purchased with his personal credit card. Keep it classified as Postage; it is not a vendor payment or a "
        "credit-card payoff."
    ),
    (
        "If the original charges are already recorded, apply the Zelle payment against the amount owed to Benjamin "
        "or the reimbursement-clearing account. Do not record Postage twice."
    ): (
        "Retain the Pirate Ship receipts and Yaakov's card statement with the reimbursement record. Do not also "
        "record the same label charges separately."
    ),
    "The rewards journals increase Other Income in March and April, not July.": (
        "The March and April rewards journals reclassify income from General Income to Credit Card Rewards; they do "
        "not change total profit or affect July."
    ),
    (
        "A July net loss of approximately $2,600 was observed. Do not attribute it to the Amazon card payoff or the "
        "capitalized machines without evidence. Rerun the monthly profit-and-loss report after all corrections, then "
        "compare July with adjacent months by revenue and account-level expense changes."
    ): (
        "The corrected July 2026 Profit and Loss report shows net profit of $367.25 on both cash and accrual bases. "
        "Compared with June, sales decreased by $10,335.75 and operating expenses increased by $5,008.90. The largest "
        "expense increases were Salaries and Employee Wages ($3,302.21) and Utilities ($2,763.60)."
    ),
}

PARAGRAPHS_TO_REMOVE = {
    "Whether the $28.13 and $700 Pirate Ship expenses were already recorded before the Zelle reimbursements.",
    "The final account-level explanation for July's approximately $2,600 loss after rerunning the corrected report.",
}

TABLE_TREATMENT_REPLACEMENTS = {
    "Rewards credit - March": (
        "Reclassified from General Income to Credit Card Rewards (Other Income); no change to total profit."
    ),
    "Rewards credit - April": (
        "Reclassified from General Income to Credit Card Rewards (Other Income); no change to total profit."
    ),
    "Total reviewed rewards credits": (
        "Credit Card Rewards (Other Income), not sales revenue; the reclassification did not change total profit, "
        "and tax treatment remains a CPA-review item."
    ),
}

NEW_CONFIRMED_ROWS: tuple[tuple[str, str, str], ...] = ()

EXPECTED_TABLE_COUNT = 1
EXPECTED_COLUMN_COUNT = 3


def set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra_paragraph in cell.paragraphs[1:]:
        remove_paragraph(extra_paragraph)


def append_formatted_row(table, values: tuple[str, str, str]) -> None:
    row_xml = deepcopy(table.rows[-1]._tr)
    table._tbl.append(row_xml)
    row = table.rows[-1]
    if len(row.cells) != EXPECTED_COLUMN_COUNT:
        raise ValueError(f"Expected {EXPECTED_COLUMN_COUNT} columns, found {len(row.cells)}")
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def update_document(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)
    found_replacements: set[str] = set()
    found_removals: set[str] = set()

    for paragraph in list(document.paragraphs):
        original_text = paragraph.text
        if original_text in PARAGRAPH_REPLACEMENTS:
            set_paragraph_text(paragraph, PARAGRAPH_REPLACEMENTS[original_text])
            found_replacements.add(original_text)
        elif original_text in PARAGRAPHS_TO_REMOVE:
            remove_paragraph(paragraph)
            found_removals.add(original_text)

    missing_replacements = set(PARAGRAPH_REPLACEMENTS) - found_replacements
    missing_removals = PARAGRAPHS_TO_REMOVE - found_removals
    if missing_replacements or missing_removals:
        raise ValueError(
            f"Missing paragraph targets: replacements={sorted(missing_replacements)}, "
            f"removals={sorted(missing_removals)}"
        )

    if len(document.tables) != EXPECTED_TABLE_COUNT:
        raise ValueError(f"Expected {EXPECTED_TABLE_COUNT} table, found {len(document.tables)}")

    table = document.tables[0]
    found_table_items: set[str] = set()
    for row in table.rows:
        item = row.cells[0].text.strip()
        if item in TABLE_TREATMENT_REPLACEMENTS:
            set_cell_text(row.cells[2], TABLE_TREATMENT_REPLACEMENTS[item])
            found_table_items.add(item)

    missing_table_items = set(TABLE_TREATMENT_REPLACEMENTS) - found_table_items
    if missing_table_items:
        raise ValueError(f"Missing table items: {sorted(missing_table_items)}")

    for values in NEW_CONFIRMED_ROWS:
        append_formatted_row(table, values)

    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    update_document(args.input, args.output)


if __name__ == "__main__":
    main()
